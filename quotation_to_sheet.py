
"""
Step 2 of 2 - read the JSON from extract_quotation.py, write the job row
into the Excel sheet, and print the material requisition.

    python quotation_to_sheet.py 1942.json

No API key, no PDF handling, no cost. Safe to run as often as you like.

Setup:
    pip install openpyxl
"""

import datetime
import json
import pathlib
import shutil
import sys
import openpyxl

# ---------------------------------------------------------------- settings

SHEET_IN = "C:/Users/jawad/OneDrive/Desktop/Fitout jobs sheet latest 22-12-2025.xlsx"
SHEET_OUT = "test_sheet.xlsx"      # set equal to SHEET_IN once you trust it
TAB = "Sheet1"

# 1-based column numbers in the sheet
COL = {
    "job": 1,
    "customer": 2,
    "project": 3,
    "qref": 4,
    "date": 5,
    "value": 7,
    "payment": 8,
}


# ------------------------------------------------------------ plain rules
# None of this is the model's job. Rules that must behave identically every
# time belong in code.

def clean_qty(value):
    """Quantities arrive as '02' or 'Lot'. Numbers become ints, Lot stays text."""
    text = str(value).strip()
    return int(text) if text.isdigit() else text


def supply_lines(line_items):
    """The material requisition is every BOQ row whose description starts
    with 'Supply of'. Relocation, installation and testing are labour."""
    return [
        item
        for item in line_items or []
        if str(item.get("description", "")).strip().lower().startswith("supply of")
    ]


def check(data):
    """Refuse to write a row that is obviously incomplete. An empty cell is
    recoverable; a wrong one buried in 2,300 rows is not."""
    problems = []
    for key in ("customer_name", "project_name", "quotation_ref", "value", "payment_term"):
        if data.get(key) in (None, ""):
            problems.append(f"{key} is empty")
    if data.get("payment_term") not in ("advance", "LPO", "credit"):
        problems.append(f"unexpected payment_term: {data.get('payment_term')!r}")
    if not data.get("line_items"):
        problems.append("no line items found")
    return problems


# ----------------------------------------------------------- sheet writing

def find_existing(ws, qref):
    """Return the row where this quotation ref already appears, if any.

    Any step with a side effect must be safe to run twice. Retries, double
    clicks and re-triggered watchers are normal, not exceptional.
    """
    target = str(qref).strip()
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, COL["qref"]).value
        if cell not in (None, "") and str(cell).strip() == target:
            return r
    return None


def write_row(data, src, dst, tab):
    """Write the job into the row below the last one that has a customer name.

    Scanning from the bottom matters: the sheet has a dozen historic gaps
    where a job number exists with no customer, and 'first empty row' would
    land in one of them.

    Returns (row, job_no), or (None, existing_row) if already processed.
    """
    if src != dst:
        shutil.copy(src, dst)

    wb = openpyxl.load_workbook(dst)
    ws = wb[tab]

    existing = find_existing(ws, data["quotation_ref"])
    if existing is not None:
        return None, existing

    last = max(
        r for r in range(2, ws.max_row + 1)
        if ws.cell(r, COL["customer"]).value not in (None, "")
    )
    row = last + 1

    job_no = ws.cell(row, COL["job"]).value       # already pre-filled in the sheet

    qref = data["quotation_ref"]
    ws.cell(row, COL["customer"]).value = data["customer_name"]
    ws.cell(row, COL["project"]).value = data["project_name"]
    ws.cell(row, COL["qref"]).value = int(qref) if str(qref).isdigit() else qref
    ws.cell(row, COL["date"]).value = datetime.date.today()   # the day payment landed
    ws.cell(row, COL["value"]).value = data["value"]
    ws.cell(row, COL["payment"]).value = data["payment_term"]

    wb.save(dst)
    return row, job_no


def print_requisition(data, job_no):
    print("\n=== material requisition ===")
    print(f"Job Number     {job_no}")
    print(f"Customer Name  {data['customer_name']}")
    print(f"Project        {data['project_name']}")
    print(f"Date           {datetime.date.today():%d/%m/%y}")
    print()
    for n, item in enumerate(supply_lines(data.get("line_items")), start=1):
        desc = str(item.get("description"))[:44]
        print(f"  {n}.  {desc:44}  qty {clean_qty(item.get('qty'))}")


# ------------------------------------------------------------------- main

if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit("usage: python quotation_to_sheet.py <json file>")

    path = pathlib.Path(sys.argv[1])
    if not path.exists():
        raise SystemExit(f"{path} not found. Run extract_quotation.py first.")

    data = json.loads(path.read_text(encoding="utf-8"))

    print("=== read from", path.name, "===")
    for key in ("customer_name", "project_name", "quotation_ref", "value", "payment_term"):
        print(f"{key:15} = {data.get(key)!r}")

    problems = check(data)
    if problems:
        print("\nNOT WRITING - the extraction looks wrong:")
        for p in problems:
            print("  -", p)
        raise SystemExit(1)

    row, job_no = write_row(data, SHEET_IN, SHEET_OUT, TAB)

    if row is None:
        print(f"\nAlready processed: quotation {data['quotation_ref']} is at row {job_no}.")
        print("Nothing written.")
        raise SystemExit(0)

    print(f"\nwrote job {job_no} to row {row} of {SHEET_OUT}")
    print_requisition(data, job_no)